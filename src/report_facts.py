"""M81: činjenice iz STVARNOG financijskog izvješća (EHO/ZSE XLSX).

Zašto postoji: analiza koja gleda samo naše agregate ne vidi ono što u
izvješću ima ime. Incident 16.09.2026. (SNBA): agenti su raspravljali o
"nerazloženih ~23,5 mil. €" dobiti, a u revidiranom GFI-ju to je stavka
"Ostali prihodi iz redovnog poslovanja" (23,72 mil. €) u godini u kojoj
je banka PRVI PUT konsolidirala stečenu štedionicu — dakle jednokratni
učinak stjecanja, a usporedba s prethodnom (nekonsolidiranom) godinom
uspoređivala je različite opsege.

Alat namjerno NE tumači: vraća što u izvješću piše (oznaka konsolidacije,
ovisni subjekti, revizor) i koje su se stavke najviše promijenile. Tumačenje
je na analitičaru/agentu, ali sada ga radi NAD ČINJENICAMA, ne nad prazninom.

Obrazac: GFI/TFI XLSX s EHO-a (listovi "Opći podaci", "Bilanca", "RDG").
Za neprepoznat raspored vraća ono što nađe — nikad ne izmišlja.
"""
from __future__ import annotations

import io
import os
import pathlib
import re
import urllib.request
from typing import Any

CACHE = pathlib.Path(os.environ.get("REPORT_CACHE", ".cache/reports"))

PDF_HINTS = ("stjecanj", "pripajanj", "preuzimanj", "poslovno spajanje",
             "povoljne kupnje", "povoljna kupnja", "negativni goodwill",
             "badwill", "konsolidir", "ovisno društvo", "ovisnog društva")

# stavke čiji skok najčešće znači jednokratni događaj (stjecanje, prodaja,
# otpis) — uvijek se izvlače eksplicitno, i kad nisu među najvećim promjenama
WATCH = (
    "ostali prihodi iz redovnog poslovanja",
    "ostali rashodi iz redovnog poslovanja",
    "prihodi na osnovi kamata",
    "rashodi na osnovi kamata",
    "opći administrativni rashodi",
    "dobit ili gubitak tekuće godine",
    "dobit ili gubitak prije oporezivanja",
    "ukupna imovina",
    "ukupni prihodi",
    "ukupni rashodi",
)


def _cells(row) -> list:
    return [c for c in row if c is not None and str(c).strip()]


def _nums(cells) -> list[float]:
    return [float(c) for c in cells if isinstance(c, (int, float))]


def _text(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip()


def kind(blob: bytes) -> str:
    """Što je zapravo preuzeto: 'xlsx', 'pdf', 'esef' (iXBRL paket) ili 'zip'.

    EHO pod istim tipom dokumenta nudi GFI obrazac (XLSX), skenirani PDF i —
    od 2025. sve češće — ESEF paket (zip u zipu, xhtml + XBRL). Bez ovoga je
    ESEF paket rušio čitanje golim KeyErrorom umjesto da se prizna."""
    import zipfile
    if blob[:4] == b"%PDF":
        return "pdf"
    if blob[:2] != b"PK":
        return "nepoznato"
    try:
        names = zipfile.ZipFile(io.BytesIO(blob)).namelist()
    except Exception:  # noqa: BLE001
        return "nepoznato"
    if "[Content_Types].xml" in names:
        return "xlsx"
    return "esef" if _esef_members(blob) else "zip"


def _esef_members(blob: bytes, depth: int = 2) -> bool:
    import zipfile
    try:
        zf = zipfile.ZipFile(io.BytesIO(blob))
        names = zf.namelist()
    except Exception:  # noqa: BLE001
        return False
    if any(n.lower().endswith((".xhtml", ".xbrl")) for n in names):
        return True
    if depth <= 0:
        return False
    return any(_esef_members(zf.read(n), depth - 1)
               for n in names if n.lower().endswith(".zip"))


def unwrap(blob: bytes, depth: int = 2) -> bytes | None:
    """Nađi GFI obrazac (XLSX) unutar zip omota; None ako ga nema."""
    import zipfile
    k = kind(blob)
    if k == "xlsx":
        return blob
    if k not in ("zip", "esef") or depth <= 0:
        return None
    zf = zipfile.ZipFile(io.BytesIO(blob))
    for n in zf.namelist():
        if n.lower().endswith(".xlsx"):
            return zf.read(n)
    for n in zf.namelist():
        if n.lower().endswith(".zip"):
            inner = unwrap(zf.read(n), depth - 1)
            if inner:
                return inner
    return None


def load_workbook(source: str | bytes):
    """Otvori XLSX s puta ili iz bajtova (bez mrežnog poziva unutar modula)."""
    import openpyxl
    if isinstance(source, bytes):
        blob = unwrap(source)
        if blob is None:
            raise ValueError(f"nije GFI obrazac (XLSX), nego: {kind(source)}")
        return openpyxl.load_workbook(io.BytesIO(blob), data_only=True)
    return openpyxl.load_workbook(source, data_only=True)


def general_info(wb) -> dict:
    """Konsolidiranost, revizija, ovisni subjekti, zaposleni — 'Opći podaci'."""
    out: dict[str, Any] = {"consolidated": None, "audited": None,
                           "subsidiaries": [], "employees": None,
                           "auditor": None, "issuer": None}
    name = next((n for n in wb.sheetnames if "opći" in n.lower()
                 or "opci" in n.lower()), None)
    if not name:
        return out
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=False))
    for i, row in enumerate(rows):
        label = _text(row[0].value if row else "")
        low = label.lower()
        # oznaka je u prvoj ćeliji DESNO od labele; H/I su legenda obrasca
        val = next((_text(c.value) for c in row[1:4] if _text(c.value)), "")
        if low.startswith("konsolidirani izvještaj"):
            out["consolidated"] = (val.upper() == "KD") if val.upper() in ("KD", "KN") else None
            out["consolidated_mark"] = val.upper() or None
        elif low.startswith("revidirano"):
            out["audited"] = (val.upper() == "RD") if val.upper() in ("RD", "RN") else None
        elif "tvrtka izdavatelja" in low:
            out["issuer"] = val or None
        elif "revizorsko društvo" in low:
            out["auditor"] = val or None
        elif "broj zaposlenih" in low:
            n = _nums(_cells([c.value for c in row]))
            if n:
                out["employees"] = int(n[-1])
        elif "ovisnih subjekata" in low:
            # popis kreće ispod naslovnog retka i traje dok ima imena
            for nxt in rows[i + 1:i + 12]:
                cells = [_text(c.value) for c in nxt if _text(c.value)]
                if not cells:
                    continue
                first = cells[0]
                # naziv subjekta nikad ne završava dvotočkom — sljedeća labela
                # obrasca zatvara popis, ma kako se zvala
                if first.endswith(":") or first.lower().startswith(
                        ("knjigovodstveni", "osoba za kontakt", "telefon",
                         "adresa", "revizorsko", "da", "ne")):
                    break
                out["subsidiaries"].append(
                    {"name": first,
                     "seat": cells[1] if len(cells) > 1 else None})
    return out


def statement_rows(wb, sheet_hint: str) -> list[dict]:
    """Retci izvještaja kao {label, prev, curr}. U GFI obrascu stupci idu
    [AOP, prethodno razdoblje, tekuće razdoblje] — uzimaju se ZADNJE dvije
    brojke retka, pa raspored s dodatnim stupcem ne ruši čitanje."""
    name = next((n for n in wb.sheetnames if sheet_hint.lower() in n.lower()), None)
    if not name:
        return []
    out = []
    for row in wb[name].iter_rows(values_only=True):
        cells = _cells(row)
        if len(cells) < 2:
            continue
        label = _text(cells[0])
        nums = _nums(cells)
        if not label or len(nums) < 2:
            continue
        prev, curr = nums[-2], nums[-1]
        out.append({"label": label, "prev": prev, "curr": curr})
    return out


def biggest_moves(rows: list[dict], top: int = 8,
                  min_abs: float = 100_000.0) -> list[dict]:
    """Stavke s najvećom APSOLUTNOM promjenom — ondje se kriju jednokratni
    događaji. Prag odbacuje šum malih pozicija."""
    moved = [{**r, "delta": r["curr"] - r["prev"]} for r in rows
             if abs(r["curr"] - r["prev"]) >= min_abs]
    moved.sort(key=lambda r: abs(r["delta"]), reverse=True)
    return moved[:top]


def watch_items(rows: list[dict]) -> list[dict]:
    out = []
    for r in rows:
        low = r["label"].lower()
        if any(w in low for w in WATCH):
            out.append({**r, "delta": r["curr"] - r["prev"]})
    return out


def facts(source: str | bytes, *, source_url: str | None = None,
          fiscal_year: int | None = None) -> dict:
    """Sve činjenice iz jednog izvješća, spremne za data_snapshot."""
    wb = load_workbook(source)
    info = general_info(wb)
    pnl = statement_rows(wb, "RDG") or statement_rows(wb, "RDG_")
    bs = statement_rows(wb, "Bilanca")
    total_assets = next((r for r in bs
                         if r["label"].lower().startswith("ukupna imovina")), None)
    return {
        "source_url": source_url,
        "fiscal_year": fiscal_year,
        "sheets": wb.sheetnames,
        "consolidated": info.get("consolidated"),
        "consolidated_mark": info.get("consolidated_mark"),
        "audited": info.get("audited"),
        "auditor": info.get("auditor"),
        "employees": info.get("employees"),
        "subsidiaries": info.get("subsidiaries"),
        "total_assets": ({"prev": total_assets["prev"], "curr": total_assets["curr"]}
                         if total_assets else None),
        "watch": watch_items(pnl),
        "biggest_moves": biggest_moves(pnl),
    }


def annual_sources(cur, ticker: str, limit: int = 2) -> list[tuple[int, str]]:
    """Zadnje `limit` fiskalnih godina s URL-om godišnjeg izvješća; unutar
    godine XLSX ima prednost pred PDF-om (strojno čitljiv obrazac)."""
    cur.execute(
        """SELECT f.fiscal_year, f.source_url
           FROM filings f JOIN companies c ON c.id = f.company_id
           WHERE c.ticker = %s AND f.doc_type = 'financial_report'
             AND f.period_type = 'annual' AND f.source_url LIKE 'http%%'
           ORDER BY f.fiscal_year DESC,
                    (f.source_url ILIKE '%%.xlsx') DESC, f.published_at DESC""",
        (ticker,))
    best: dict[int, str] = {}
    for fy, url in cur.fetchall():
        best.setdefault(fy, url)
    return sorted(best.items(), key=lambda kv: kv[0], reverse=True)[:limit]


def fetch(url: str) -> bytes | None:
    """Dohvat izvješća uz lokalni cache (ista datoteka se ne skida dvaput)."""
    CACHE.mkdir(parents=True, exist_ok=True)
    key = CACHE / url.rsplit("/", 1)[-1][:120]
    if key.exists():
        return key.read_bytes()
    try:
        with urllib.request.urlopen(url, timeout=90) as r:
            data = r.read()
        key.write_bytes(data)
        return data
    except Exception:  # noqa: BLE001 — pozivatelj bilježi 'nije dohvaćeno'
        return None


def text_hints(blob: bytes) -> list[str]:
    """Izvješće koje nije GFI obrazac (PDF ili ESEF paket) — izvlači se tekst
    i traže naznake jednokratnih/statusnih događaja koje analitičar MORA
    pročitati. Prazan popis znači 'nije nađeno', ne 'nema ga'."""
    k = kind(blob)
    if k == "pdf":
        return pdf_scan(blob)
    if k in ("esef", "zip"):
        return _esef_scan(blob)
    return []


def _esef_scan(blob: bytes, depth: int = 2) -> list[str]:
    """ESEF/iXBRL paket: xhtml je tekst, pa se naznake traže izravno u njemu."""
    import zipfile
    try:
        zf = zipfile.ZipFile(io.BytesIO(blob))
        names = zf.namelist()
    except Exception:  # noqa: BLE001
        return []
    found: set[str] = set()
    for n in names:
        low = n.lower()
        if low.endswith((".xhtml", ".html", ".htm")):
            try:
                txt = zf.read(n)[:40_000_000].decode("utf-8", "ignore").lower()
            except Exception:  # noqa: BLE001
                continue
            found |= {h for h in PDF_HINTS if h in txt}
        elif low.endswith(".zip") and depth > 0:
            try:
                found |= set(_esef_scan(zf.read(n), depth - 1))
            except Exception:  # noqa: BLE001
                continue
    return sorted(found)


def pdf_scan(blob: bytes) -> list[str]:
    """PDF izvješće nije strojno čitljiv obrazac — izvlači se tekst i traže
    naznake jednokratnih/statusnih događaja koje analitičar MORA pročitati."""
    try:
        import pymupdf
    except Exception:  # noqa: BLE001
        try:
            import fitz as pymupdf  # starije izdanje paketa
        except Exception:  # noqa: BLE001
            return []
    try:
        with pymupdf.open(stream=blob, filetype="pdf") as doc:
            text = "\n".join(p.get_text() for p in doc[:40]).lower()
    except Exception:  # noqa: BLE001
        return []
    return sorted({h for h in PDF_HINTS if h in text})


def consolidation_changed(prev_facts: dict, curr_facts: dict) -> bool:
    """TRUE kad je izvještaj prešao iz nekonsolidiranog u konsolidirani (ili
    obrnuto) — tada postotne usporedbe dviju godina mjere različite opsege i
    NE SMIJU se prikazivati bez upozorenja."""
    a, b = prev_facts.get("consolidated"), curr_facts.get("consolidated")
    return a is not None and b is not None and a != b
